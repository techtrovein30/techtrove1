import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os
import datetime

def set_cell_background(cell, hex_color):
    """Sets background color of a cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding for a cell (in twips)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles & Colors
    COLOR_PRIMARY = RGBColor(109, 40, 217)   # Deep Purple / TechTrove accent
    COLOR_DARK = RGBColor(30, 41, 59)        # Slate 800
    COLOR_MUTED = RGBColor(100, 116, 139)    # Slate 500
    COLOR_HIGH = RGBColor(220, 38, 38)       # Red 600
    COLOR_MID = RGBColor(217, 119, 6)        # Amber 600
    COLOR_LOW = RGBColor(37, 99, 235)        # Blue 600

    # Document Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_sub = title_p.add_run("TECHTROVE 3.0 SYMPOSIUM PLATFORM\n")
    run_sub.font.size = Pt(11)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_PRIMARY

    run_title = title_p.add_run("Security Audit, Bug Report & Comprehensive Code Review")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_DARK

    # Metadata Block
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(18)
    meta_run = meta_p.add_run(f"Audit Date: {datetime.date.today().strftime('%B %d, %Y')}   |   Status: Comprehensive Assessment   |   Target: TechTrove 3.0 Production Stack")
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = COLOR_MUTED

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ─── SECTION 1: EXECUTIVE SUMMARY ───────────────────────────────────────
    h1 = doc.add_heading("1. Executive Summary", level=1)
    h1.style.font.color.rgb = COLOR_PRIMARY

    p_exec = doc.add_paragraph(
        "A rigorous, full-stack architectural security audit and code review was performed on the TechTrove 3.0 codebase. "
        "The application manages symposium registrations, participant authentication, team rosters, multi-tier payments, "
        "and administrative workflows across PostgreSQL (Supabase), React/TypeScript, and Storage APIs. "
        "This assessment evaluates data integrity, financial security, authentication vectors, performance bottlenecks, and certificate issuance requirements."
    )
    p_exec.paragraph_format.space_after = Pt(12)

    # Summary Statistics Table
    summary_table = doc.add_table(rows=1, cols=4)
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary_table.autofit = False

    col_widths = [Inches(1.6), Inches(1.6), Inches(1.6), Inches(1.6)]
    headers = ["Total Issues", "High Severity", "Mid Severity", "Low Severity"]
    hdr_cells = summary_table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(10)
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], "4C1D95") # Dark Purple
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        hdr_cells[i].width = col_widths[i]

    row_cells = summary_table.add_row().cells
    stats = [("18", COLOR_DARK), ("4", COLOR_HIGH), ("8", COLOR_MID), ("6", COLOR_LOW)]
    for i, (val, color) in enumerate(stats):
        row_cells[i].text = val
        row_cells[i].paragraphs[0].runs[0].font.bold = True
        row_cells[i].paragraphs[0].runs[0].font.size = Pt(18)
        row_cells[i].paragraphs[0].runs[0].font.color.rgb = color
        row_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_background(row_cells[i], "F8FAFC")
        row_cells[i].width = col_widths[i]

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # ─── SECTION 2: SEVERITY DEFINITIONS ────────────────────────────────────
    h2_def = doc.add_heading("2. Severity Classification Matrix", level=1)
    h2_def.style.font.color.rgb = COLOR_PRIMARY

    sev_table = doc.add_table(rows=4, cols=3)
    sev_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sev_widths = [Inches(1.2), Inches(1.8), Inches(3.4)]

    s_headers = ["Severity", "Impact Level", "Definition & Criteria"]
    for i, h in enumerate(s_headers):
        c = sev_table.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(c, "1E293B")
        c.width = sev_widths[i]

    sev_data = [
        ("HIGH", "Critical / Exploit", "Direct data loss, payment evasion, CSV formula injection, RLS bypasses, silent cascading deletions, or core database architectural blockades (e.g. certificate generation block)."),
        ("MID", "Moderate / Functional", "Inconsistent reporting metrics, race conditions, false pending statuses for free events, unindexed slow queries, or thundering-herd realtime refetch cascades."),
        ("LOW", "Minor / Quality", "MIME-type spoofing bypasses, leftover merge conflict artifacts, bundle bloat (>500 kB warning), unused stub functions, or missing type hints.")
    ]

    bg_colors = ["FEE2E2", "FEF3C7", "DBEAFE"]
    fg_colors = [COLOR_HIGH, COLOR_MID, COLOR_LOW]

    for idx, (sev, imp, desc) in enumerate(sev_data):
        row = sev_table.rows[idx + 1].cells
        row[0].text = sev
        row[0].paragraphs[0].runs[0].font.bold = True
        row[0].paragraphs[0].runs[0].font.color.rgb = fg_colors[idx]
        set_cell_background(row[0], bg_colors[idx])
        row[0].width = sev_widths[0]

        row[1].text = imp
        row[1].paragraphs[0].runs[0].font.bold = True
        row[1].paragraphs[0].runs[0].font.size = Pt(9.5)
        row[1].width = sev_widths[1]

        row[2].text = desc
        row[2].paragraphs[0].runs[0].font.size = Pt(9.5)
        row[2].width = sev_widths[2]

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # ─── SECTION 3: DETAILED AUDIT FINDINGS ─────────────────────────────────
    h3 = doc.add_heading("3. Detailed Vulnerability & Bug Findings", level=1)
    h3.style.font.color.rgb = COLOR_PRIMARY

    issues = [
        # ── HIGH ISSUES ─────────────────────────────────────────────────────
        {
            "id": "SEC-01",
            "title": "CSV Formula / Command Injection (CWE-1236) in Admin Exports",
            "severity": "HIGH",
            "component": "AdminRegistrationsPage.tsx / AdminPaymentsPage.tsx / AdminTeamsPage.tsx",
            "description": "The admin export function concatenates raw input into CSV strings using straightforward template literals without neutralizing formula execution characters (=, +, -, @, tab, CR). Furthermore, quotes are not properly RFC-4180 escaped.",
            "impact": "If an attacker registers with a team name or captain name such as '=cmd|/C calc'!A0 or '=HYPERLINK(...)', Microsoft Excel or Google Sheets executes arbitrary system commands or exfiltrates confidential spreadsheet data when an administrator downloads and opens the registrant CSV roster.",
            "remediation": "Import and strictly utilize a hardened RFC-4180 sanitizer (e.g. csv.ts) that prepends single quotes (') to any field beginning with formula characters (=, +, -, @, \\t, \\r, \\n) and doubles internal quotation marks."
        },
        {
            "id": "DAT-02",
            "title": "Team Member JSON Encapsulation & Certificate Generation Blockade",
            "severity": "HIGH",
            "component": "Database Schema (registrations_internal / registrations_external)",
            "description": "All team members and substitute players are packed into a single 'members jsonb' column. Individual participants do not possess dedicated database rows, unique foreign-key IDs, or indexable columns.",
            "impact": "Certificates cannot be cleanly generated, addressed to individual unique IDs, or verified via public QR codes. Furthermore, tracking individual attendance or querying a student's symposium participation across multiple teams requires expensive, unindexed JSON table scans.",
            "remediation": "Deploy migration 20240101000006_registration_members.sql to establish a dedicated 'registration_members' table populated automatically by an AFTER INSERT/UPDATE PostgreSQL trigger. This maintains single-request frontend registration while providing clean individual rows for certificate issuance and verification."
        },
        {
            "id": "DAT-03",
            "title": "Unprotected Cascade Deletion of Events Destroys Registrations & Payments",
            "severity": "HIGH",
            "component": "supabase/migrations/20240101000000_initial_schema.sql / eventStore.ts",
            "description": "The foreign key in registrations_internal and registrations_external references public.events(id) with ON DELETE CASCADE. When an admin deletes an event via adminDeleteEvent(eventId), PostgreSQL silently purges all student registrations, rosters, and financial audit records for that event.",
            "impact": "Accidental deletion of an event from the Admin Events page instantly eradicates all associated student bookings and payment verifications with no recovery path.",
            "remediation": "Alter the foreign key to ON DELETE RESTRICT (or NO ACTION). Prevent event deletion if registrations exist, prompting administrators to archive or cancel the event instead."
        },
        {
            "id": "SEC-04",
            "title": "Unrestricted Arbitrary Storage URL Acceptance in getUploadSignedUrl",
            "severity": "HIGH",
            "component": "src/lib/storage.ts (Line 148)",
            "description": "The getUploadSignedUrl helper checks if the stored path begins with 'http://' or 'https://' and immediately returns it without verifying origin domain or hostname.",
            "impact": "An attacker can submit an external phishing URL, malware payload, or malicious tracking link in place of a valid Supabase storage path, which is then rendered directly in admin payment proof modals.",
            "remediation": "Validate that any absolute URL strictly originates from the project's designated Supabase storage domain (e.g., regex check against the project host), and reject untrusted external schemes."
        },

        # ── MID ISSUES ──────────────────────────────────────────────────────
        {
            "id": "FIN-05",
            "title": "Inconsistent Financial Metric Aggregation Across Admin Views",
            "severity": "MID",
            "component": "src/pages/admin/AdminPaymentsPage.tsx vs src/lib/adminApi.ts (getAdminStats)",
            "description": "In AdminPaymentsPage.tsx, registrations with paymentStatus === 'confirmed' (such as free internal registrations) fall into the else block and increment pendingCount and pendingRevenue. Conversely, in adminApi.ts, they increment recorded.",
            "impact": "Administrators see contradictory statistics: the main Dashboard reports 0 pending revenue for free internal students, while the Payments screen reports them as pending collections.",
            "remediation": "Standardize payment status evaluation across all admin pages. Treat 'confirmed' and 'recorded' as verified non-pending statuses, and only flag 'pending' for external registrations awaiting review."
        },
        {
            "id": "UI-06",
            "title": "False 'Pending' Payment Badge Displayed to Free Internal Students",
            "severity": "MID",
            "component": "src/pages/ProfilePage.tsx (Line 32 - RegStatusBadge)",
            "description": "The RegStatusBadge component in ProfilePage.tsx only checks 'status === \"recorded\"' to display the green 'Paid' badge. All other statuses default to an amber 'Pending' badge.",
            "impact": "Internal SIMATS students who legitimately registered for free receive a 'Pending' badge in amber on their profile, creating confusion and prompting unnecessary support queries.",
            "remediation": "Update RegStatusBadge to treat 'confirmed' and 'recorded' as green badges ('Confirmed' / 'Free' / 'Paid') and reserve amber strictly for 'pending'."
        },
        {
            "id": "SEC-07",
            "title": "Public Registration Lookup by Code Exposes Participant PII",
            "severity": "MID",
            "component": "src/lib/api.ts (getRegistrationByCode)",
            "description": "The API endpoint getRegistrationByCode(code) returns complete registration objects (including all member full names, phone numbers, registration numbers, emails, and payment screenshots) without verifying caller authorization.",
            "impact": "Because registration codes follow a predictable format (TT-XXXXXX), unauthorized actors could brute-force or scrape codes to harvest student contact details and ID numbers.",
            "remediation": "Enforce ownership verification (auth.uid() === user_id) or mask sensitive member phone numbers and emails on public receipt views."
        },
        {
            "id": "REL-08",
            "title": "Thundering Herd Full-Table Rescans on Supabase Realtime Events",
            "severity": "MID",
            "component": "src/lib/useAdminRealtime.ts / AdminDashboardPage.tsx",
            "description": "The realtime synchronization hooks listen for any wildcard postgres_changes on internal and external registration tables and invoke full table fetch operations (adminListRegistrations / fetchStats).",
            "impact": "During heavy registration periods, 100 simultaneous student sign-ups trigger 100 full table scans across the entire database, leading to network congestion, rate limiting, and client lag.",
            "remediation": "Debounce the realtime refresh handler (e.g. 1.5 second throttle) or apply incremental state patching from payload.new instead of executing full-table re-queries."
        },
        {
            "id": "AUT-09",
            "title": "Client-Side Username Collision Loop and TOCTOU Race Condition",
            "severity": "MID",
            "component": "src/context/AuthContext.tsx (completeGoogleProfile)",
            "description": "Username uniqueness is checked via an iterative client-side loop executing up to 50 sequential checks across both participant tables before inserting.",
            "impact": "If two users with identical names complete onboarding simultaneously, both check the database at the same instant (Time-of-Check to Time-of-Use) and attempt to claim the same username, causing constraint crashes. Additionally, making up to 100 roundtrips on slow networks degrades UX.",
            "remediation": "Move username allocation into a PostgreSQL trigger or server RPC using gen_random_uuid() suffix or sequence numbers."
        },
        {
            "id": "FIN-10",
            "title": "Lack of Unique Constraint on External Payment UTR Numbers",
            "severity": "MID",
            "component": "supabase/migrations/20240101000000_initial_schema.sql",
            "description": "The utr_number column on registrations_external lacks a unique database constraint or duplicate-detection check.",
            "impact": "Malicious or confused external participants can submit the same banking UTR number across multiple team registrations, enabling fee evasion unless manually caught by an admin.",
            "remediation": "Add a unique index or trigger verification on non-null utr_number values within registrations_external."
        },
        {
            "id": "DAT-11",
            "title": "Multi-Event Fee Pass Trigger Desynchronization with Client Code",
            "severity": "MID",
            "component": "supabase/migrations/20240101000005_tech_pass.sql vs src/lib/api.ts",
            "description": "The 20240101000005_tech_pass.sql migration expects multi-event transactions to share the identical registration_code. However, api.ts generates a new random registration code per call.",
            "impact": "If the client submits multiple events as separate createRegistration calls, each event receives a different code, causing the database to bill ₹75 on every single event instead of applying the flat pass.",
            "remediation": "Ensure the batch registration payload explicitly supplies a shared transaction/pass code or updates api.ts to handle event arrays atomically."
        },
        {
            "id": "CRY-12",
            "title": "Weak PRNG (Math.random) for Critical Registration and Event Identifiers",
            "severity": "MID",
            "component": "src/lib/api.ts (makeId) / src/lib/eventStore.ts (makeEventId)",
            "description": "makeId relies on Math.random().toString(36).slice(2, 8). Math.random is cryptographically insecure, predictable, and prone to birthday paradox collisions.",
            "impact": "When code collisions occur on unique database columns (registration_code, id), the insert fails with unhandled PostgreSQL 23505 unique_violation errors, failing the user's booking.",
            "remediation": "Use crypto.getRandomValues() with alphanumeric alphabets and add a single-retry loop on collision."
        },

        # ── LOW ISSUES ──────────────────────────────────────────────────────
        {
            "id": "MNT-13",
            "title": "Orphaned Git Merge Conflict Artifact in Repository",
            "severity": "LOW",
            "component": "registerpage_conflict.txt",
            "description": "A raw git merge conflict snippet file was inadvertently committed into the repository root.",
            "impact": "Repo clutter and potential developer confusion regarding current production registration logic.",
            "remediation": "Delete registerpage_conflict.txt from version control."
        },
        {
            "id": "SEC-14",
            "title": "Client-Side MIME-Type Header Spoofing on File Uploads",
            "severity": "LOW",
            "component": "src/lib/storage.ts (validateUploadFile)",
            "description": "File format validation checks file.type, which is populated by the browser strictly based on file extension rather than binary magic byte inspection.",
            "impact": "Non-image files renamed with .png extensions can be uploaded to the private bucket, though execution is prevented by Supabase storage sandboxing.",
            "remediation": "Verify image file signatures (magic bytes) client-side before upload or enforce server-side image transcoding."
        },
        {
            "id": "PRD-15",
            "title": "Large Production Bundle Chunk Warning (>500 kB)",
            "severity": "LOW",
            "component": "vite.config.ts / dist/assets/index.js (595 kB)",
            "description": "Vite production build emits warnings that chunks exceed 500 kB due to monlithic vendor packaging of Lucide icons, Supabase, and routing dependencies.",
            "impact": "Increased First Contentful Paint (FCP) and Time-to-Interactive (TTI) on low-bandwidth mobile networks.",
            "remediation": "Configure manualChunks in vite.config.ts for vendor libraries and implement React.lazy for Admin routes."
        },
        {
            "id": "COD-16",
            "title": "Hardcoded Static Days Fallback in Event Store",
            "severity": "LOW",
            "component": "src/lib/eventStore.ts (Line 217)",
            "description": "getDays() falls back to staticDays while background fetch is running, which can cause brief UI rendering flickering if database event schedules differ from static arrays.",
            "impact": "Transient layout shift during initial cold app launch.",
            "remediation": "Exclusively rely on the asynchronous useAllEvents hook with proper skeleton loaders."
        },
        {
            "id": "COD-17",
            "title": "Dead / Deprecated Admin Authentication Helper Functions",
            "severity": "LOW",
            "component": "src/lib/adminApi.ts (isCurrentUserAdmin / seedAdminIfNeeded)",
            "description": "isCurrentUserAdmin() always returns false for backward compatibility, and seedAdminIfNeeded() is an empty no-op.",
            "impact": "Code bloat and minor confusion for external developers maintaining the project.",
            "remediation": "Remove deprecated stub functions and update call sites to use AuthContext user.role."
        },
        {
            "id": "UX-18",
            "title": "Missing Empty State Feedback on Filtered Admin Roster Views",
            "severity": "LOW",
            "component": "src/pages/admin/AdminTeamsPage.tsx",
            "description": "When filters match 0 teams, the accordion view collapses completely without an explanatory 'No matching teams' card.",
            "impact": "Admin users may perceive the application as frozen or broken when a filter returns 0 records.",
            "remediation": "Add an explicit empty state placeholder with clear search reset action."
        }
    ]

    for item in issues:
        # Issue Title & Badge
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(12)
        p_title.paragraph_format.space_after = Pt(4)
        p_title.paragraph_format.keep_with_next = True

        run_id = p_title.add_run(f"[{item['id']}] ")
        run_id.font.bold = True
        run_id.font.size = Pt(11)
        run_id.font.color.rgb = COLOR_DARK

        run_name = p_title.add_run(f"{item['title']} ")
        run_name.font.bold = True
        run_name.font.size = Pt(12)
        run_name.font.color.rgb = COLOR_DARK

        # Severity Badge
        sev_color = COLOR_HIGH if item["severity"] == "HIGH" else (COLOR_MID if item["severity"] == "MID" else COLOR_LOW)
        run_sev = p_title.add_run(f"({item['severity']} SEVERITY)")
        run_sev.font.bold = True
        run_sev.font.size = Pt(9.5)
        run_sev.font.color.rgb = sev_color

        # Details Table
        tbl = doc.add_table(rows=4, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        t_widths = [Inches(1.5), Inches(5.1)]

        fields = [
            ("Affected Area", item["component"]),
            ("Description", item["description"]),
            ("Risk & Impact", item["impact"]),
            ("Remediation", item["remediation"]),
        ]

        for r_idx, (label, val) in enumerate(fields):
            c_lbl = tbl.rows[r_idx].cells[0]
            c_lbl.text = label
            c_lbl.paragraphs[0].runs[0].font.bold = True
            c_lbl.paragraphs[0].runs[0].font.size = Pt(9)
            c_lbl.paragraphs[0].runs[0].font.color.rgb = COLOR_MUTED
            set_cell_background(c_lbl, "F1F5F9")
            c_lbl.width = t_widths[0]

            c_val = tbl.rows[r_idx].cells[1]
            c_val.text = val
            c_val.paragraphs[0].runs[0].font.size = Pt(9.5)
            c_val.width = t_widths[1]

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ─── SECTION 4: REMEDIATION ROADMAP ─────────────────────────────────────
    h4 = doc.add_heading("4. Prioritized Remediation Roadmap", level=1)
    h4.style.font.color.rgb = COLOR_PRIMARY

    p_road = doc.add_paragraph()
    p_road.add_run("To ensure absolute stability, security, and smooth symposium execution, address findings according to the following phased schedule:\n\n")

    phases = [
        ("Phase 1: Critical Hotfixes (Immediate / Day 1)", [
            "Apply CSV injection defenses across all admin export routines.",
            "Deploy migration 20240101000006_registration_members.sql to establish dedicated certificate and member tables.",
            "Sanitize storage URLs in getUploadSignedUrl to eliminate arbitrary URL reflection.",
            "Add ON DELETE RESTRICT on events to prevent catastrophic registration wipes."
        ]),
        ("Phase 2: Data & Financial Accuracy (Day 2 - 3)", [
            "Align AdminPaymentsPage and Dashboard metrics so confirmed free registrations are never reported as pending.",
            "Fix ProfilePage RegStatusBadge so internal participants see 'Confirmed' rather than 'Pending'.",
            "Debounce realtime admin hooks to eliminate thundering herd database spikes.",
            "Add unique constraints on external UTR payment entries."
        ]),
        ("Phase 3: Codebase Cleanliness & Optimization (Day 4+)", [
            "Configure Vite manualChunks to split oversized bundles under 500 kB.",
            "Remove dead stub functions in adminApi.ts and delete registerpage_conflict.txt.",
            "Implement client-side magic byte inspection for file uploads."
        ])
    ]

    for phase_title, steps in phases:
        p_phase = doc.add_paragraph()
        p_phase.paragraph_format.space_before = Pt(8)
        p_phase.paragraph_format.space_after = Pt(2)
        r_ph = p_phase.add_run(phase_title)
        r_ph.font.bold = True
        r_ph.font.size = Pt(11)
        r_ph.font.color.rgb = COLOR_PRIMARY

        for s in steps:
            p_step = doc.add_paragraph(style='List Bullet')
            p_step.paragraph_format.space_before = Pt(0)
            p_step.paragraph_format.space_after = Pt(2)
            r_s = p_step.add_run(s)
            r_s.font.size = Pt(9.5)

    # Output file
    output_filename = "s:/techtrove1/TechTrove_Bug_Report_and_Code_Review.docx"
    doc.save(output_filename)
    print(f"Successfully generated: {output_filename}")

if __name__ == "__main__":
    create_report()
